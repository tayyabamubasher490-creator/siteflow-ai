from __future__ import annotations

import csv
import json
import os
from pathlib import Path
from typing import Any

from crewai.tools import tool
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


def _read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        pages.append(f"\n--- PDF page {index + 1} ---\n{text}")
    return "\n".join(pages)


def _read_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _read_xlsx(path: Path) -> str:
    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    chunks: list[str] = []
    for sheet in workbook.worksheets:
        chunks.append(f"\n--- Excel sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(value.strip() for value in values):
                chunks.append(" | ".join(values))
    return "\n".join(chunks)


def _read_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        rows = csv.reader(file)
        return "\n".join(" | ".join(row) for row in rows)


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


@tool("Read project documents")
def read_project_documents(directory: str) -> str:
    """Read all supported project/payment documents from a directory.

    Use this tool before making conclusions about uploaded evidence.
    Supported formats: PDF, DOCX, XLSX, CSV, TXT, MD.
    """
    root = Path(directory)
    if not root.exists():
        return f"ERROR: directory does not exist: {directory}"

    supported = {".pdf", ".docx", ".xlsx", ".csv", ".txt", ".md"}
    files = sorted(p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in supported)

    if not files:
        return "ERROR: no supported documents were found."

    output: list[str] = []
    total_chars = 0
    max_chars = 90000

    for path in files:
        try:
            suffix = path.suffix.lower()
            if suffix == ".pdf":
                text = _read_pdf(path)
            elif suffix == ".docx":
                text = _read_docx(path)
            elif suffix == ".xlsx":
                text = _read_xlsx(path)
            elif suffix == ".csv":
                text = _read_csv(path)
            else:
                text = _read_text(path)

            text = text.strip()
            block = f"\n\n===== FILE: {path.name} =====\n{text}\n"
            remaining = max_chars - total_chars
            if remaining <= 0:
                break
            output.append(block[:remaining])
            total_chars += len(block)
        except Exception as exc:
            output.append(f"\n===== FILE: {path.name} =====\nREAD ERROR: {exc}\n")

    return "".join(output)


@tool("Calculate payment amount")
def calculate_payment_amount(
    gross_value: float,
    previous_certified: float = 0.0,
    retention_percent: float = 0.0,
    other_deductions: float = 0.0,
) -> str:
    """Calculate current gross, retention, deductions and net payment.

    Use this tool for payment arithmetic rather than estimating amounts manually.
    """
    values = [gross_value, previous_certified, retention_percent, other_deductions]
    if any(value < 0 for value in values):
        return "ERROR: monetary inputs and percentages cannot be negative."

    if retention_percent > 100:
        return "ERROR: retention_percent cannot exceed 100."

    current_before_retention = max(gross_value - previous_certified, 0.0)
    retention = current_before_retention * (retention_percent / 100.0)
    net = current_before_retention - retention - other_deductions

    result = {
        "gross_value": round(gross_value, 2),
        "previous_certified": round(previous_certified, 2),
        "current_before_retention": round(current_before_retention, 2),
        "retention_percent": retention_percent,
        "retention_amount": round(retention, 2),
        "other_deductions": round(other_deductions, 2),
        "net_payment": round(net, 2),
        "warning": "Net payment is negative." if net < 0 else None,
    }
    return json.dumps(result, indent=2)


@tool("Check required documents")
def check_required_documents(
    available_document_names: str,
    required_document_names: str,
) -> str:
    """Compare available document names with required document names.

    Inputs are comma-separated names. Matching is case-insensitive and partial.
    """
    available = [
        item.strip().lower()
        for item in available_document_names.split(",")
        if item.strip()
    ]
    required = [
        item.strip()
        for item in required_document_names.split(",")
        if item.strip()
    ]

    missing: list[str] = []
    found: list[str] = []

    for requirement in required:
        requirement_lower = requirement.lower()
        if any(requirement_lower in item or item in requirement_lower for item in available):
            found.append(requirement)
        else:
            missing.append(requirement)

    return json.dumps({"found": found, "missing": missing}, indent=2)


@tool("Record audit event")
def record_audit_event(directory: str, agent_name: str, event: str) -> str:
    """Append a timestamped audit event to the current run's audit_log.json."""
    from datetime import datetime, timezone

    root = Path(directory)
    root.mkdir(parents=True, exist_ok=True)
    log_path = root / "audit_log.json"

    try:
        existing: list[dict[str, Any]] = (
            json.loads(log_path.read_text(encoding="utf-8"))
            if log_path.exists()
            else []
        )
    except Exception:
        existing = []

    existing.append(
        {
            "timestamp_utc": datetime.now(timezone.utc).isoformat(),
            "agent": agent_name,
            "event": event[:1000],
        }
    )
    log_path.write_text(json.dumps(existing, indent=2), encoding="utf-8")
    return f"Audit event recorded for {agent_name}."
